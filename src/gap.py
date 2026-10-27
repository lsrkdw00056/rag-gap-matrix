import copy
import json
import os
import re
import time
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from google import genai
from google.genai import types
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()

EXTRACTION = Path("output/extraction.json")
CORRECTIONS = Path("output/manual_corrections.json")
META_CSV = Path("papers_metadata.csv")
OUT_XLSX = Path("output/gap_matrix.xlsx")
OUT_MD = Path("output/gap_analysis.md")
LLM_MODEL = "gemini-3.1-flash-lite"

# Fokus usulan paper kelompok. Ubah sesuai keputusan kelompok.
FOKUS_USULAN = (
    "sistem RAG kesehatan berbahasa Indonesia dengan evaluasi faithfulness "
    "per klaim dan perbandingan sistematis strategi chunking"
)

LABELS = {
    "masalah_tujuan": "Masalah / Tujuan",
    "subdomain_kesehatan": "Sub-domain Kesehatan",
    "retriever_embedding": "Retriever / Embedding",
    "generator_llm": "Generator (LLM)",
    "strategi_chunking": "Strategi Chunking",
    "dataset_knowledge_base": "Dataset / Knowledge Base",
    "metrik_evaluasi": "Metrik Evaluasi",
    "evaluasi_faithfulness": "Evaluasi Faithfulness / Halusinasi",
    "bahasa_data": "Bahasa Data",
    "hasil_utama": "Hasil Utama",
    "keterbatasan": "Keterbatasan",
    "future_work": "Future Work",
}

GAP_PROMPT = """You are helping a student group write a literature gap analysis on
Retrieval-Augmented Generation (RAG) in the healthcare domain.
Below is a matrix extracted from {n} papers and then checked by humans.

Cell status values:
- "Not reported": the paper does not state it (a real gap candidate).
- "Partial": only an indirect proxy is reported (for example similarity to a reference answer).
- "N/A": not applicable to this paper type (for example a survey has no chunking). N/A is NOT a gap.

{table}

Statistics per column (applicable papers exclude N/A):
{stats}

Operational definition: a paper measures faithfulness or hallucination explicitly ONLY if it reports
a metric that compares the answer with the retrieved evidence or source documents, or a hallucination rate.
Reference-similarity metrics (BLEU, METEOR, cosine to a reference answer) count as "Partial".
Statements that only describe mitigation count as "Not reported".

Rules:
- Use ONLY the matrix. Do not invent facts.
- Cite a paper ID only if its cell directly supports the claim.
- Do not write "most studies" unless more than half of the applicable papers support it.
- Surveys and conceptual papers must not be described as implemented systems.
- Papers that have a domain note are not healthcare-specific; do not use them as evidence for healthcare-specific claims.

Write in Indonesian.
1. Identify 4 to 6 research gaps. For each give: "gap" (short title), "bukti" (which paper IDs report or omit what),
   "alasan" (why it matters), "usulan" (a concrete research idea that could fill it).
2. Write "analisis_paragraf": exactly 3 paragraphs in formal academic Indonesian.
   Paragraph 1: what existing studies already do (cite paper IDs like (P01, P06)).
   Paragraph 2: the main gaps found, with evidence from the matrix.
   Paragraph 3: the proposed group paper, focusing on: {focus}. State how it differs from the closest existing study in the matrix.

Return ONLY JSON: {{"gaps": [{{"gap": "", "bukti": "", "alasan": "", "usulan": ""}}], "analisis_paragraf": ""}}
"""


def is_na(v) -> bool:
    return str(v).strip().lower().startswith("n/a")


def is_partial(v) -> bool:
    return str(v).strip().lower().startswith("partial")


def is_missing(v) -> bool:
    s = str(v).strip().lower()
    return s.startswith("not reported") or s in ("", "error")


def cell_text(ans: dict) -> str:
    v = str(ans.get("value", "")).strip()
    pages = ans.get("evidence_pages") or []
    if is_missing(v) or not pages:
        return v
    return f"{v} (hal. {', '.join(str(p) for p in pages)})"


def acc_label(correct: int, checked: int, belum: int) -> str:
    if not checked:
        return "-"
    label = f"{correct / checked * 100:.0f}%"
    return label if belum == 0 else label + " (belum final)"


def parse_json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(),
                  flags=re.MULTILINE).strip()
    m = re.search(r"\{.*\}", text, re.DOTALL)
    return json.loads(m.group(0) if m else text)


def ask_llm(client, prompt: str):
    cfg = types.GenerateContentConfig(response_mime_type="application/json")
    for attempt in range(4):
        try:
            resp = client.models.generate_content(
                model=LLM_MODEL, contents=prompt, config=cfg)
            return parse_json(resp.text)
        except Exception as e:
            wait = 20 * (attempt + 1)
            print(
                f"   ! percobaan {attempt + 1} gagal ({str(e)[:150]}), tunggu {wait}s")
            time.sleep(wait)
    return None


def style_sheet(ws, width=40, color_status=False):
    head_fill = PatternFill("solid", fgColor="1F3864")
    miss_fill = PatternFill("solid", fgColor="F8CBAD")   # Not reported
    part_fill = PatternFill("solid", fgColor="FFE699")   # Partial
    na_fill = PatternFill("solid", fgColor="D9D9D9")     # N/A
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if color_status:
                t = str(cell.value or "").strip().lower()
                if t.startswith("not reported"):
                    cell.fill = miss_fill
                elif t.startswith("partial"):
                    cell.fill = part_fill
                elif t.startswith("n/a"):
                    cell.fill = na_fill
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "B2"


def main():
    ext = json.loads(EXTRACTION.read_text(encoding="utf-8"))
    raw = copy.deepcopy(ext)  # simpan hasil asli LLM untuk perbandingan

    corr = json.loads(CORRECTIONS.read_text(encoding="utf-8")
                      ) if CORRECTIONS.exists() else {}
    corrections = corr.get("corrections", {})
    notes = corr.get("catatan_paper", {})
    validated = corr.get("validated_papers", [])
    verified = corr.get("verified_ok", {})

    # terapkan koreksi manual
    corrected = {}  # (paper, kolom) -> jenis ("salah" atau "standardisasi")
    for pid, cols in corrections.items():
        for col, new in cols.items():
            if pid not in ext or col not in LABELS:
                print(
                    f"! koreksi diabaikan (paper/kolom tidak dikenal): {pid} / {col}")
                continue
            base = ext[pid].get(col, {})
            ext[pid][col] = {
                **base,
                "value": new["value"],
                "evidence_pages": new.get("evidence_pages", []),
                "confidence": "manual",
            }
            corrected[(pid, col)] = new.get("jenis", "salah")
    print(f"{len(corrected)} sel dikoreksi manual dari {OUT_MD.parent / 'manual_corrections.json'}")

    meta = {}
    if META_CSV.exists():
        df = pd.read_csv(META_CSV).fillna("")
        meta = {r["paper_id"]: r.to_dict() for _, r in df.iterrows()}

    rows = []
    missing = {l: [] for l in LABELS.values()}
    partial = {l: [] for l in LABELS.values()}
    na = {l: [] for l in LABELS.values()}
    for pid in sorted(ext):
        m = meta.get(pid, {})
        row = {
            "Paper": pid,
            "Penulis": m.get("penulis", ""),
            "Tahun": m.get("tahun", ""),
            "Judul": m.get("judul", ""),
        }
        for col, label in LABELS.items():
            ans = ext[pid].get(col, {"value": "ERROR"})
            row[label] = cell_text(ans)
            v = ans.get("value", "")
            if is_na(v):
                na[label].append(pid)
            elif is_partial(v):
                partial[label].append(pid)
            elif is_missing(v):
                missing[label].append(pid)
        row["Catatan Domain"] = notes.get(pid, "")
        rows.append(row)
    matrix = pd.DataFrame(rows)
    n = len(rows)

    stat_rows = []
    for label in LABELS.values():
        applicable = n - len(na[label])
        stat_rows.append(
            {
                "Kolom": label,
                "Berlaku": applicable,
                "Not reported": len(missing[label]),
                "Partial": len(partial[label]),
                "N/A": len(na[label]),
                "Persen Not reported": f"{len(missing[label]) / applicable * 100:.0f}%" if applicable else "-",
                "Paper Not reported": ", ".join(missing[label]) or "-",
                "Paper Partial": ", ".join(partial[label]) or "-",
            }
        )
    stats = pd.DataFrame(stat_rows)
    print(stats[["Kolom", "Berlaku", "Not reported", "Partial",
          "N/A", "Persen Not reported"]].to_string(index=False))

    # tabel koreksi dan validasi
    koreksi_cols = ["Paper", "Kolom", "Jenis", "Nilai LLM",
                    "Nilai Koreksi", "Alasan", "Dihitung di akurasi"]
    koreksi_rows = []
    for (pid, col), jenis in sorted(corrected.items()):
        koreksi_rows.append(
            {
                "Paper": pid,
                "Kolom": LABELS[col],
                "Jenis": "Kesalahan ekstraksi" if jenis == "salah" else "Standardisasi",
                "Nilai LLM": cell_text(raw[pid].get(col, {})),
                "Nilai Koreksi": cell_text(ext[pid][col]),
                "Alasan": corrections[pid][col].get("alasan", ""),
                "Dihitung di akurasi": "Ya" if pid in validated else "Tidak",
            }
        )
    koreksi_df = pd.DataFrame(koreksi_rows, columns=koreksi_cols)

    val_cols = ["Paper", "Sel dicek", "Benar",
                "Salah", "Belum dicek", "Akurasi"]
    val_rows, tot_checked, tot_correct, tot_wrong = [], 0, 0, 0
    for pid in validated:
        ok = {c for c in verified.get(pid, []) if c in LABELS}
        fixed = {c for (p, c) in corrected if p == pid}
        wrong = {c for (p, c), j in corrected.items()
                 if p == pid and j == "salah"}
        checked = ok | fixed
        correct = checked - wrong
        belum = [c for c in LABELS if c not in checked]
        val_rows.append(
            {
                "Paper": pid,
                "Sel dicek": len(checked),
                "Benar": len(correct),
                "Salah": len(wrong),
                "Belum dicek": len(belum),
                "Akurasi": acc_label(len(correct), len(checked), len(belum)),
            }
        )
        tot_checked += len(checked)
        tot_correct += len(correct)
        tot_wrong += len(wrong)
        if belum:
            print(f"! {pid}: {len(belum)} sel belum dicek, akurasi belum final")
    if val_rows:
        val_rows.append(
            {
                "Paper": "TOTAL",
                "Sel dicek": tot_checked,
                "Benar": tot_correct,
                "Salah": tot_wrong,
                "Belum dicek": sum(r["Belum dicek"] for r in val_rows),
                "Akurasi": acc_label(tot_correct, tot_checked, sum(r["Belum dicek"] for r in val_rows)),
            }
        )
    validasi_df = pd.DataFrame(val_rows, columns=val_cols)

    # ringkasan untuk LLM
    lines = []
    for r in rows:
        lines.append(f"## {r['Paper']} ({r['Penulis']}, {r['Tahun']})")
        if r["Catatan Domain"]:
            lines.append(f"- Domain note: {r['Catatan Domain']}")
        for label in LABELS.values():
            lines.append(f"- {label}: {str(r[label])[:350]}")
    stats_text = "\n".join(
        f"- {s['Kolom']}: Not reported {s['Not reported']}/{s['Berlaku']} ({s['Paper Not reported']}); partial: {s['Paper Partial']}"
        for _, s in stats.iterrows()
    )

    client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
    print("\nMenyusun kandidat gap dengan LLM ...")
    out = ask_llm(
        client,
        GAP_PROMPT.format(n=n, table="\n".join(lines),
                          stats=stats_text, focus=FOKUS_USULAN),
    )

    gaps, analisis = [
    ], "(Gagal menyusun analisis. Jalankan ulang python src/gap.py)"
    if out:
        gaps = out.get("gaps", [])
        analisis = out.get("analisis_paragraf", analisis)
    gaps_df = pd.DataFrame(gaps) if gaps else pd.DataFrame(
        columns=["gap", "bukti", "alasan", "usulan"])
    gaps_df = gaps_df.map(lambda x: ", ".join(
        map(str, x)) if isinstance(x, list) else x)
    gaps_df.columns = [c.capitalize() for c in gaps_df.columns]

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as xw:
        matrix.to_excel(xw, sheet_name="Gap Matrix", index=False)
        stats.to_excel(xw, sheet_name="Statistik Gap", index=False)
        gaps_df.to_excel(xw, sheet_name="Kandidat Gap", index=False)
        pd.DataFrame({"Analisis Gap": analisis.split("\n\n")}).to_excel(
            xw, sheet_name="Analisis", index=False)
        koreksi_df.to_excel(xw, sheet_name="Koreksi Manual", index=False)
        validasi_df.to_excel(xw, sheet_name="Validasi", index=False)

        ws = xw.sheets["Gap Matrix"]
        style_sheet(ws, 40, color_status=True)
        # sel yang dikoreksi manual
        blue = PatternFill("solid", fgColor="BDD7EE")
        paper_order, col_order = list(matrix["Paper"]), list(matrix.columns)
        for pid, col in corrected:
            ws.cell(row=paper_order.index(pid) + 2,
                    column=col_order.index(LABELS[col]) + 1).fill = blue
        style_sheet(xw.sheets["Statistik Gap"], 28)
        style_sheet(xw.sheets["Kandidat Gap"], 50)
        style_sheet(xw.sheets["Analisis"], 140)
        style_sheet(xw.sheets["Koreksi Manual"], 45)
        style_sheet(xw.sheets["Validasi"], 18)

    md = ["# Analisis Gap\n", analisis, "\n## Kandidat Gap\n"]
    for i, g in enumerate(gaps, 1):
        md.append(
            f"{i}. **{g.get('gap', '')}**\n   - Bukti: {g.get('bukti', '')}\n"
            f"   - Alasan: {g.get('alasan', '')}\n   - Usulan: {g.get('usulan', '')}\n"
        )
    OUT_MD.write_text("\n".join(md), encoding="utf-8")
    print(f"\nSelesai. File: {OUT_XLSX} dan {OUT_MD}")


if __name__ == "__main__":
    main()
